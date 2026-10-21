#!/usr/bin/env python3
"""Download verified public inputs and rebuild the PRJNA422760 input tables.

Inputs are downloaded from their public sources and verified before analysis.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


BASE_DIR = Path(__file__).resolve().parent
GITHUB_COMMIT = "61ec285181d82dfa273d1cabbc4e0b72a00fee9c"
RDS_URL = (
    "https://raw.githubusercontent.com/lynchlab-ucsf/SinusitisMicrobiome/"
    f"{GITHUB_COMMIT}/AnalysisPhyloseq.rds"
)
RDS_SHA256 = "249bc53778863e739fd5039c47d6c46c589d0d2cbe44dd0cc236419c4afe8fc6"
WORKBOOK_URL = (
    "https://pmc.ncbi.nlm.nih.gov/articles/instance/6372837/bin/"
    "mSystems.00169-18-st003.xlsx"
)
WORKBOOK_ARTICLE_URL = "https://pmc.ncbi.nlm.nih.gov/articles/PMC6372837/"
WORKBOOK_SHA256 = "7c8aa33043e31d4c7de701ae0c4aa6c4de5d81a6b93055a26b39b12e5f5694c4"
RUNINFO_URL = "https://trace.ncbi.nlm.nih.gov/Traces/sra-db-be/runinfo?acc=PRJNA422760"
USER_AGENT = "airway-microbiome-reproducibility/1.1.0"

BASELINE_SHA256 = {
    "data/processed/prjna422760_sample_metadata.tsv": (
        "1e293196699799fe992846ce353511f29e52f49e4d995f148a5274fd99274d2d"
    ),
    "data/processed/prjna422760_genus_abundance.tsv": (
        "fa6b3e181b7aa780e8318bf1e66270f44f39aa0e0033536a809d300413fb6961"
    ),
}


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _curl_bytes(url: str, *, referer: str | None, cookie: str | None, timeout: int) -> bytes:
    command = [
        "curl",
        "--fail",
        "--location",
        "--silent",
        "--show-error",
        "--max-time",
        str(timeout),
        "-H",
        f"User-Agent: {USER_AGENT}",
    ]
    if referer:
        command.extend(["-H", f"Referer: {referer}"])
    if cookie:
        command.extend(["-H", f"Cookie: {cookie}"])
    command.append(url)
    result = subprocess.run(command, capture_output=True, check=False)
    if result.returncode:
        detail = result.stderr.decode("utf-8", errors="replace").strip()
        raise RuntimeError(f"curl failed for {url}: {detail or result.returncode}")
    return result.stdout


def request_bytes(
    url: str,
    *,
    referer: str | None = None,
    cookie: str | None = None,
    timeout: int = 90,
) -> bytes:
    headers = {"User-Agent": USER_AGENT}
    if referer:
        headers["Referer"] = referer
    if cookie:
        headers["Cookie"] = cookie
    request = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = response.read()
            if response.status < 200 or response.status >= 300:
                raise RuntimeError(f"HTTP {response.status} for {url}")
            return payload
    except (urllib.error.URLError, TimeoutError, OSError, RuntimeError):
        # Some networks terminate Python TLS sessions while allowing curl.
        # Both paths return bytes for validation below; neither writes directly.
        return _curl_bytes(url, referer=referer, cookie=cookie, timeout=timeout)


def atomic_write(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".part")
    try:
        temporary.write_bytes(payload)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def download_pinned_rds(path: Path) -> dict[str, Any]:
    payload = request_bytes(RDS_URL)
    digest = sha256_bytes(payload)
    if digest != RDS_SHA256:
        raise ValueError(f"AnalysisPhyloseq.rds SHA-256 mismatch: {digest}")
    if not payload.startswith(b"\x1f\x8b"):
        raise ValueError("AnalysisPhyloseq.rds failed gzip/RDS signature validation")
    atomic_write(path, payload)
    return {
        "artifact": "AnalysisPhyloseq.rds",
        "kind": "public_source",
        "source_url": RDS_URL,
        "output_path": path,
        "sha256": digest,
        "expected_sha256": RDS_SHA256,
        "bytes": len(payload),
        "provenance": f"lynchlab-ucsf/SinusitisMicrobiome at commit {GITHUB_COMMIT}",
    }


def solve_cloudpmc_pow(page: bytes) -> str | None:
    text = page.decode("utf-8", errors="replace")
    challenge_match = re.search(r"POW_CHALLENGE\s*=\s*['\"]([^'\"]+)['\"]", text)
    difficulty_match = re.search(r"POW_DIFFICULTY\s*=\s*['\"]?(\d+)['\"]?", text)
    if not challenge_match or not difficulty_match:
        return None
    challenge = challenge_match.group(1)
    difficulty = int(difficulty_match.group(1))
    if difficulty < 0 or difficulty > 8:
        raise ValueError(f"Unexpected CloudPMC proof-of-work difficulty: {difficulty}")
    prefix = "0" * difficulty
    nonce = 0
    while True:
        digest = hashlib.sha256(f"{challenge}{nonce}".encode("utf-8")).hexdigest()
        if digest.startswith(prefix):
            return f"cloudpmc-viewer-pow={challenge},{nonce}"
        nonce += 1


def validate_xlsx(payload: bytes) -> None:
    if not payload.startswith(b"PK\x03\x04"):
        raise ValueError("Supplement download is not an XLSX ZIP file; refusing to save it")
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise ValueError(f"Corrupt XLSX member: {bad_member}")
        if "xl/workbook.xml" not in archive.namelist():
            raise ValueError("ZIP payload has no XLSX workbook.xml; refusing to save it")


def download_pinned_workbook(path: Path) -> dict[str, Any]:
    payload = request_bytes(WORKBOOK_URL, referer=WORKBOOK_ARTICLE_URL)
    if not payload.startswith(b"PK\x03\x04"):
        cookie = solve_cloudpmc_pow(payload)
        if cookie is None:
            raise ValueError("PMC returned non-XLSX content without a solvable public challenge")
        payload = request_bytes(WORKBOOK_URL, referer=WORKBOOK_ARTICLE_URL, cookie=cookie)
    validate_xlsx(payload)
    digest = sha256_bytes(payload)
    if digest != WORKBOOK_SHA256:
        raise ValueError(f"PRJNA422760 workbook SHA-256 mismatch: {digest}")
    atomic_write(path, payload)
    return {
        "artifact": "mSystems.00169-18-st003.xlsx",
        "kind": "public_source",
        "source_url": WORKBOOK_URL,
        "output_path": path,
        "sha256": digest,
        "expected_sha256": WORKBOOK_SHA256,
        "bytes": len(payload),
        "provenance": "PMC6372837 supplementary table; Genus worksheet",
    }


def download_runinfo(path: Path, sample_names: set[str]) -> tuple[pd.DataFrame, dict[str, Any]]:
    payload = request_bytes(RUNINFO_URL)
    try:
        text = payload.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text)))
    except (UnicodeDecodeError, csv.Error) as exc:
        raise ValueError("NCBI RunInfo response is not readable CSV") from exc
    required = {"Run", "BioSample", "SampleName", "Model", "LibraryStrategy"}
    if not rows or not required.issubset(rows[0]):
        raise ValueError("NCBI RunInfo response is missing required public metadata fields")
    if any(str(row.get("Run", "")).lower().startswith("<!doctype") for row in rows):
        raise ValueError("NCBI RunInfo returned HTML; refusing to save it as metadata")
    metadata = pd.DataFrame(rows, dtype=str)
    if metadata["Run"].duplicated().any() or metadata["SampleName"].duplicated().any():
        raise ValueError("NCBI RunInfo contains duplicated run or sample names")
    if set(metadata["SampleName"]) != sample_names:
        missing = sorted(sample_names - set(metadata["SampleName"]))
        extra = sorted(set(metadata["SampleName"]) - sample_names)
        raise ValueError(f"NCBI RunInfo samples do not match the supplement (missing={missing}, extra={extra})")
    if len(metadata) != 56:
        raise ValueError(f"Expected 56 PRJNA422760 sequencing runs; received {len(metadata)}")
    atomic_write(path, payload)
    return metadata, {
        "artifact": "PRJNA422760_runinfo.csv",
        "kind": "public_source",
        "source_url": RUNINFO_URL,
        "output_path": path,
        "sha256": sha256_bytes(payload),
        "expected_sha256": "recorded at build time; NCBI RunInfo is a live endpoint",
        "bytes": len(payload),
        "provenance": "NCBI SRA RunInfo; run, BioSample, sample name and sequencing platform metadata",
    }


def merged_label_for_column(worksheet: Any, row: int, column: int) -> str:
    value = worksheet.cell(row=row, column=column).value
    if value is not None:
        return str(value)
    for merged_range in worksheet.merged_cells.ranges:
        if (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col <= column <= merged_range.max_col
        ):
            anchor = worksheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
            return str(anchor) if anchor is not None else ""
    return ""


def extract_tables(workbook_path: Path, runinfo: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=True)
    if "Genus" not in workbook.sheetnames:
        raise ValueError("The public supplement does not contain a Genus worksheet")
    worksheet = workbook["Genus"]
    sample_columns: list[dict[str, Any]] = []
    for column in range(2, worksheet.max_column + 1):
        sample_id = worksheet.cell(row=4, column=column).value
        if sample_id is None:
            continue
        sample_id = str(sample_id).strip()
        if not re.fullmatch(r"\d+[ac]", sample_id):
            continue
        group_label = merged_label_for_column(worksheet, 3, column)
        if group_label.startswith("HP"):
            clinical_group = "HP"
        elif group_label.startswith("SOM"):
            clinical_group = "SOM"
        else:
            raise ValueError(f"Sample {sample_id} has no published HP/SOM column-group label")
        tissue = "adenoid" if sample_id.endswith("a") else "tonsil"
        patient_id = sample_id[:-1]
        sample_columns.append(
            {
                "column": column,
                "sample_id": sample_id,
                "sample_name": f"{sample_id}_{tissue}",
                "patient_id": patient_id,
                "tissue": tissue,
                "clinical_group": clinical_group,
                "published_group_label": group_label,
            }
        )

    if len(sample_columns) != 56:
        raise ValueError(f"Expected 56 workbook sample columns; found {len(sample_columns)}")
    sample_ids = [row["sample_id"] for row in sample_columns]
    if len(set(sample_ids)) != 56:
        raise ValueError("The Genus worksheet contains duplicate sample IDs")
    patients: dict[str, list[dict[str, Any]]] = {}
    for sample in sample_columns:
        patients.setdefault(sample["patient_id"], []).append(sample)
    if len(patients) != 28 or any(
        sorted(row["tissue"] for row in pair) != ["adenoid", "tonsil"]
        or len({row["clinical_group"] for row in pair}) != 1
        for pair in patients.values()
    ):
        raise ValueError("The supplement does not encode 28 complete, group-consistent tissue pairs")

    run_metadata = runinfo[["Run", "BioSample", "SampleName", "Model", "LibraryStrategy"]].rename(
        columns={"Run": "run", "BioSample": "biosample", "SampleName": "sample_name", "Model": "instrument_model"}
    )
    sample_metadata = pd.DataFrame(sample_columns).drop(columns=["column"])
    sample_metadata = sample_metadata.merge(run_metadata, on="sample_name", how="left", validate="one_to_one")
    if sample_metadata[["run", "biosample", "instrument_model", "LibraryStrategy"]].isna().any().any():
        raise ValueError("Some workbook samples could not be matched to public NCBI run metadata")

    abundance_rows = []
    for row in range(5, worksheet.max_row + 1):
        genus = worksheet.cell(row=row, column=1).value
        if genus is None or not str(genus).strip():
            continue
        genus = str(genus).strip()
        for sample in sample_columns:
            value = worksheet.cell(row=row, column=sample["column"]).value
            relative_abundance = float(value or 0.0)
            if not pd.notna(relative_abundance) or relative_abundance < 0:
                raise ValueError(f"Invalid abundance for {sample['sample_name']} / {genus}")
            abundance_rows.append(
                {
                    "sample_name": sample["sample_name"],
                    "sample_id": sample["sample_id"],
                    "genus": genus,
                    "relative_abundance": relative_abundance,
                }
            )
    genus_abundance = pd.DataFrame(abundance_rows)
    if len(genus_abundance) != 4480 or genus_abundance["sample_name"].nunique() != 56:
        raise ValueError(f"Unexpected genus table shape: {len(genus_abundance)} observations")
    workbook.close()
    return sample_metadata, genus_abundance


def build_provenance() -> str:
    return """# Public input provenance

`AnalysisPhyloseq.rds` is downloaded from the `lynchlab-ucsf/SinusitisMicrobiome` repository at the fixed commit recorded in `DATA_MANIFEST.tsv`. Its SHA-256 is checked before the file is saved.

The PRJNA422760 abundance and group labels come from the public article supplement [mSystems.00169-18-st003.xlsx](https://pmc.ncbi.nlm.nih.gov/articles/instance/6372837/bin/mSystems.00169-18-st003.xlsx), `Genus` worksheet. Row 3 contains the published HP-group and SOM-group column headers, including merged cells; row 4 contains sample IDs. The builder carries those labels forward as `clinical_group` and `published_group_label`. HP denotes the hyperplasia group and SOM denotes secretory otitis media as defined in the accompanying article ([PMC6372837](https://pmc.ncbi.nlm.nih.gov/articles/PMC6372837/)). These labels are not inferred from sequencing metadata. The suffix `a` identifies adenoid tissue and `c` identifies tonsil tissue; the numeric part of each sample ID is the paired patient identifier.

Per-genus percentages are copied from the supplement without renormalization. Every source genus row is retained, including the two published Fusobacterium rows; downstream analysis sums rows with the same sample and genus before abundance tests, state assignments and distances. Public NCBI SRA RunInfo is used only to attach run, BioSample, instrument and library-strategy fields.

The downstream paired genus analysis tests nine selected genera in three overlapping strata (`all`, `HP` and `SOM`), yielding one 27-test family. Each contrast is the patient-level tonsil-minus-adenoid difference in percentage points; two-sided Wilcoxon signed-rank tests use `zero_method=wilcox`, the normal approximation, and no continuity correction, followed by one Benjamini-Hochberg adjustment across the 27 P values. The all-patient and subgroup contrasts overlap because the subgroup patients are included in the all-patient analysis. The input builder generates data only and does not run these tests.
"""


def write_manifest(path: Path, rows: list[dict[str, Any]]) -> None:
    fields = [
        "artifact",
        "kind",
        "source_url",
        "output_path",
        "sha256",
        "expected_sha256",
        "baseline_sha256",
        "bytes",
        "provenance",
    ]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            record = dict(row)
            record["output_path"] = str(Path(record["output_path"]).relative_to(path.parent))
            writer.writerow(record)


def build_output(output_dir: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    rds_path = output_dir / "data/references/sinusitis_microbiome/AnalysisPhyloseq.rds"
    workbook_path = output_dir / "data/references/external_labels/PRJNA422760/mSystems.00169-18-st003.xlsx"
    runinfo_path = output_dir / "data/raw/metadata/PRJNA422760_runinfo.csv"
    records.append(download_pinned_rds(rds_path))
    records.append(download_pinned_workbook(workbook_path))

    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=True)
    worksheet = workbook["Genus"]
    workbook_sample_names = set()
    for column in range(2, worksheet.max_column + 1):
        sample_id = worksheet.cell(row=4, column=column).value
        if sample_id is not None and re.fullmatch(r"\d+[ac]", str(sample_id).strip()):
            tissue = "adenoid" if str(sample_id).strip().endswith("a") else "tonsil"
            workbook_sample_names.add(f"{str(sample_id).strip()}_{tissue}")
    workbook.close()
    runinfo, runinfo_record = download_runinfo(runinfo_path, workbook_sample_names)
    records.append(runinfo_record)

    sample_metadata, genus_abundance = extract_tables(workbook_path, runinfo)
    derived = [
        (
            "data/processed/prjna422760_sample_metadata.tsv",
            sample_metadata,
            "Workbook sample IDs, tissue and published group labels joined to NCBI run metadata",
        ),
        (
            "data/processed/prjna422760_genus_abundance.tsv",
            genus_abundance,
            "Genus percentages transcribed row-wise from the public Genus worksheet",
        ),
    ]
    for relative_path, frame, provenance in derived:
        destination = output_dir / relative_path
        destination.parent.mkdir(parents=True, exist_ok=True)
        frame.to_csv(destination, sep="\t", index=False)
        digest = hashlib.sha256(destination.read_bytes()).hexdigest()
        expected = BASELINE_SHA256[relative_path]
        if digest != expected:
            raise ValueError(
                f"Rebuilt input differs from the reference input checksums for {relative_path}: {digest}"
            )
        records.append(
            {
                "artifact": Path(relative_path).name,
                "kind": "derived_input",
                "source_url": WORKBOOK_URL if "genus_abundance" in relative_path else f"{WORKBOOK_URL}; {RUNINFO_URL}",
                "output_path": destination,
                "sha256": digest,
                "expected_sha256": expected,
                "baseline_sha256": expected,
                "bytes": destination.stat().st_size,
                "provenance": provenance,
            }
        )

    provenance_path = output_dir / "DATA_PROVENANCE.md"
    provenance_path.write_text(build_provenance(), encoding="utf-8")
    manifest_path = output_dir / "DATA_MANIFEST.tsv"
    write_manifest(manifest_path, records)
    return records


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    records = build_output(root)
    docs = root / "docs"
    docs.mkdir(exist_ok=True)
    (root / "DATA_MANIFEST.tsv").replace(docs / "DATA_MANIFEST.tsv")
    (root / "DATA_PROVENANCE.md").replace(docs / "DATA_PROVENANCE.md")
    print("Downloaded and verified public inputs; reconstructed 56 profiles from 28 paired patients.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
